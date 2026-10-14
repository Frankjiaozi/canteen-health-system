"""生成导入模板。

模板的列顺序与 sheet 名就是导入的契约，改了这里等于改了契约，
所以模板与解析必须放在一起看。
"""

from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

#: sheet 名 → (表头, 示例行)
SHEETS: dict[str, tuple[list[str], list[Any]]] = {
    "食材": (
        [
            "名称",
            "类别",
            "能量(千卡/100g)",
            "蛋白质(克/100g)",
            "脂肪(克/100g)",
            "碳水化合物(克/100g)",
            "钠(毫克/100g)",
            "可食部比例(0-1)",
            "单价(元/千克)",
            "数据来源",
            "来源版本",
            "备注",
        ],
        [
            "西兰花",
            "蔬菜类",
            36,
            4.1,
            0.6,
            4.3,
            18.8,
            0.83,
            6.0,
            "《中国食物成分表（标准版第 6 版）》",
            "6版",
            "",
        ],
    ),
    "调味料": (
        [
            "名称",
            "类别",
            "单位",
            "能量(千卡/100g)",
            "蛋白质(克/100g)",
            "脂肪(克/100g)",
            "碳水化合物(克/100g)",
            "钠(毫克/100g)",
            "单价(元/千克)",
            "数据来源",
            "备注",
        ],
        [
            "大豆油",
            "食用油",
            "g",
            899,
            0,
            99.9,
            0,
            0,
            12.0,
            "《中国食物成分表（标准版第 6 版）》",
            "",
        ],
    ),
    "菜品": (
        ["名称", "类别", "烹饪方式", "标准份量(克)", "单锅份数", "过敏原(顿号分隔)", "备注"],
        ["清炒西兰花", "素菜", "炒", 120, 40, "", ""],
    ),
    "配方": (
        ["菜品名称", "行类型(食材/调味料)", "名称", "单份用量(克)", "可食部比例覆盖"],
        ["清炒西兰花", "食材", "西兰花", 110, ""],
    ),
}

_HEADER_FILL = PatternFill("solid", fgColor="D9E8FF")
_HEADER_FONT = Font(bold=True)


def build_template() -> bytes:
    """生成带示例行与填写说明的模板工作簿。"""
    workbook = Workbook()
    # 去掉默认建的那张空表
    workbook.remove(workbook.worksheets[0])

    for sheet_name, (headers, example) in SHEETS.items():
        sheet = workbook.create_sheet(sheet_name)
        sheet.append(headers)
        sheet.append(example)

        for index, header in enumerate(headers, start=1):
            cell = sheet.cell(row=1, column=index)
            cell.fill = _HEADER_FILL
            cell.font = _HEADER_FONT
            cell.alignment = Alignment(horizontal="center")
            # 列宽按表头长度粗略撑开，省得使用者第一眼看到一堆挤在一起的列
            sheet.column_dimensions[get_column_letter(index)].width = max(12, len(header) + 4)
        sheet.freeze_panes = "A2"

    _add_readme(workbook)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _add_readme(workbook: Workbook) -> None:
    sheet = workbook.create_sheet("填写说明", 0)
    lines = [
        "知味 · 校园智慧膳食平台 —— 批量导入模板",
        "",
        "一、四个 sheet 的含义：",
        "  食材   —— 每 100 克可食部的营养成分。数据来源必填，这是可追溯的硬要求。",
        "  调味料 —— 油、盐、糖等。营养成分也按每 100 克填，与食材库同一口径。",
        "  菜品   —— 菜品的基础信息。",
        "  配方   —— 配方明细，一行一个「食材或调味料」。同一道菜写多行。",
        "",
        "二、导入顺序与规则：",
        "  1) 先建食材与调味料，再建菜品，最后建配方——配方里引用的名称必须已经存在。",
        "  2) 名称已存在的按「更新」处理，不存在的按「新增」处理；菜品配方会被整体替换，配方版本自动加一。",
        "  3) 类别、烹饪方式可以直接写中文（蔬菜类、炒），也可以写代码（VEGETABLE、STIR_FRY）。",
        "  4) 留空的营养项表示「数据缺失」，请保持留空，不要填 0——填 0 会让营养核算静默失真。",
        "  5) 可食部比例填 0~1 之间的小数，例如鸡蛋填 0.85；留空表示 1（全部可食）。",
        "  6) 配方里的「可食部比例覆盖」只有需要覆盖食材库默认值时才填。",
        "",
        "三、导入方式：",
        "  在「菜品与配方」页面点「批量导入」，先选文件校验，确认报告没有问题后再执行导入。",
        "  校验不通过时不会有任何数据被写入。",
    ]
    for line in lines:
        sheet.append([line])
    sheet.column_dimensions["A"].width = 100
